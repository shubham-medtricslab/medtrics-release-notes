"""Tests for the /kb-articles pipeline scripts in scripts/kb/ — selection,
placement, grouping, trimmed inputs, capture helpers, the tenant screen,
annotation, spec assembly, rendering, packaging, PDF checks and usage
apportioning. Hermetic: no browser, no Mem, no network. Tests that need
python-docx, pypdf or Pillow skip when the module is missing.

Pins: non-features and ledger-done tickets dropped before any agent; placement
order (known ticket > URL prefix > unique module); multi-persona tickets get a
capture visit per persona but one writer group; groups split at the cap;
acceptance criteria split out and UI strings taken from ADDED lines only;
expiry detection and resume plans; DOM box -> image box; dom_box only when
OCR is absent and the DOM text matches exactly; the demo allowlist; the
annotation ring is orange and the watermark red; an unverified shot becomes a
missing-shot marker; label_as_written inferred from bold; prose failures split
from evidence failures; the draft render carries no provenance by default.
"""

import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from kb_fixtures import golden, has_module, write_assets

import kb_capture as C
import kb_inputs as I
import kb_select as SEL
import kb_spec as SP
import kb_usage as U
import kb_validate as V
import qa_gif_to_shots as Q

ROUTES = {
    "demo": {"base_url": "https://demo", "domain": "medtrics.sg", "login_path": "/accounts/login/",
             "allow_strings": ["Medtrics Medical Center"], "home": {}},
    "personas": {"admin": {"portal": "Administration portal", "audience": "Program Administrator"},
                 "faculty": {"portal": "Personal portal", "audience": "Faculty"},
                 "external": {"portal": "Link", "audience": "External", "no_session": True}},
    "uncapturable_statuses": ["No demo data"],
    "surfaces": [
        {"surface": "Forms > Assignments", "persona": "admin", "also_personas": ["faculty"],
         "url": "/forms/evaluation/assignments/", "recipe": [],
         "match": {"modules": ["Evaluations & Forms"], "url_prefixes": ["/forms/evaluation/assignments/"],
                   "template_globs": []}, "known_tickets": {"M1-10": "Captured"}},
        {"surface": "Gradebook > Charts", "persona": "admin", "url": "/gradebook/grouped/", "recipe": [],
         "match": {"modules": ["Gradebook"], "url_prefixes": ["/gradebook/grouped/"],
                   "template_globs": ["src/templates/gradebook/*.html"]},
         "known_tickets": {"M1-20": "No demo data"}},
        {"surface": "Faculty > Notes", "persona": "faculty", "url": "", "recipe": [],
         "match": {"modules": ["Notes"], "url_prefixes": [], "template_globs": []},
         "known_tickets": {}},
    ],
}


def issue(tid, ttype="feature", module="", diffs=(), description=""):
    return {"identifier": tid, "title": f"Title {tid}", "task_type": ttype, "module": module,
            "description": description,
            "mrs": [{"diffs": [{"new_path": p, "diff": d} for p, d in diffs]}]}


class Select(unittest.TestCase):
    def test_filter_type_then_ledger(self):
        kept, skipped = SEL.filter_issues(
            [issue("M1-1", "bug"), issue("M1-2"), issue("M1-3", "critical_feature"), issue("M1-4")],
            {"M1-2": {"status": "drafted"}, "M1-4": {"status": "partial"}})
        self.assertEqual([i["identifier"] for i in kept], ["M1-3"])
        self.assertEqual({s["ticket"]: s["reason"] for s in skipped},
                         {"M1-1": "type: bug", "M1-2": "ledger: drafted", "M1-4": "ledger: partial"})
        kept, _ = SEL.filter_issues([issue("M1-4")], {"M1-4": {"status": "partial"}},
                                    retry_partial=True)
        self.assertEqual(len(kept), 1)

    def test_filter_cli_ids_limits_release_data(self):
        with tempfile.TemporaryDirectory() as d:
            Path(d, "data.json").write_text(json.dumps(
                {"issues": [issue("M1-1"), issue("M1-2"), issue("M1-3", "bug")]}))
            import contextlib, io
            with contextlib.redirect_stdout(io.StringIO()):
                SEL.main(["filter", "--data", f"{d}/data.json", "--ids", "m1-2,M1-3",
                          "--out", f"{d}/f.json"])
            out = json.loads(Path(d, "f.json").read_text())
            self.assertEqual([i["identifier"] for i in out["issues"]], ["M1-2"])
            self.assertEqual(out["skipped"], [{"ticket": "M1-3", "reason": "type: bug"}])

    def test_no_code_skipped_not_found_listed(self):
        no_code = {"identifier": "M1-50", "title": "T", "task_type": "feature", "mrs": []}
        kept, skipped = SEL.filter_issues([no_code, issue("M1-51")], {})
        self.assertEqual([i["identifier"] for i in kept], ["M1-51"])
        self.assertEqual(skipped, [{"ticket": "M1-50", "reason": SEL.NO_CODE}])
        kept, _ = SEL.filter_issues([no_code], {}, allow_no_code=True)
        self.assertEqual(len(kept), 1)
        with tempfile.TemporaryDirectory() as d:
            Path(d, "data.json").write_text(json.dumps(
                {"issues": [issue("M1-51"), issue("M1-99")],
                 "dropped": [{"identifier": "M1-52", "task_type": "bug"}]}))
            import contextlib, io
            with contextlib.redirect_stdout(io.StringIO()):
                SEL.main(["filter", "--data", f"{d}/data.json", "--ids", "M1-51,M1-52,M1-53",
                          "--out", f"{d}/f.json"])
            out = json.loads(Path(d, "f.json").read_text())
            self.assertEqual([i["identifier"] for i in out["issues"]], ["M1-51"])
            self.assertEqual({s["ticket"]: s["reason"] for s in out["skipped"]},
                             {"M1-52": "type: bug", "M1-53": "not found in Optimus"})

    def test_place_order(self):
        s, how = SEL.place_issue(issue("M1-10"), ROUTES)
        self.assertEqual((s["surface"], how), ("Forms > Assignments", "known"))
        s, how = SEL.place_issue(issue("M1-30", diffs=[("src/x/urls.py", "+ path('gradebook/grouped/')")]),
                                 ROUTES)
        self.assertEqual((s["surface"], how), ("Gradebook > Charts", "url"))
        s, how = SEL.place_issue(issue("M1-31", diffs=[("src/templates/gradebook/a.html", "+x")]), ROUTES)
        self.assertEqual(how, "template")
        s, how = SEL.place_issue(issue("M1-32", module="Notes"), ROUTES)
        self.assertEqual((s["surface"], how), ("Faculty > Notes", "module"))
        self.assertEqual(SEL.place_issue(issue("M1-33", module="Nope"), ROUTES), (None, None))

    def test_url_rule_only_added_urls_py_lines(self):
        body = "+ x = '/gradebook/grouped/'\n- path('gradebook/grouped/', v)\n"
        for path in ("src/x/views.py", "src/x/urls.py"):
            self.assertEqual(SEL.place_issue(issue("M1-34", diffs=[(path, body)]), ROUTES), (None, None))
        s, how = SEL.place_issue(issue("M1-35", diffs=[("src/x/urls.py", "+    path('gradebook/grouped/', v)")]),
                                 ROUTES)
        self.assertEqual(how, "url")

    def test_ticket_route_placement(self):
        desc = ("**URL/Page:** https://demo.medtricslab.com/gradebook/grouped/\n\n"
                "Opens a Session > hits View Button\n")
        s, how = SEL.place_issue(issue("M1-36", description=desc), ROUTES)
        self.assertEqual(how, "ticket_route")
        p = SEL._placement(issue("M1-36", description=desc), s, how, ROUTES)
        self.assertEqual((p["persona"], p["surface"], p["url"], p["capturable"]),
                         ("admin", "/gradebook/grouped/", "/gradebook/grouped/", True))
        self.assertEqual(p["route"]["clicks"], [{"row": True}, {"text": "View"}])
        # unmatched path: admin persona, the path is the surface
        d2 = "**URL/Page:** https://demo.medtricslab.com/cases/administration/demographic-list/\n"
        s, how = SEL.place_issue(issue("M1-37", description=d2), ROUTES)
        p = SEL._placement(issue("M1-37", description=d2), s, how, ROUTES)
        self.assertEqual((p["persona"], p["surface"]), ("admin", "/cases/administration/demographic-list/"))
        # a personal-portal page ("/my/") is the faculty persona
        d4 = "**URL/Page:** https://lmucvm.medtrics.sg/forms/evaluation/my/#Review\n"
        s, how = SEL.place_issue(issue("M1-38", description=d4), ROUTES)
        p = SEL._placement(issue("M1-38", description=d4), s, how, ROUTES)
        self.assertEqual((p["persona"], p["url"]), ("faculty", "/forms/evaluation/my/#Review"))
        # a known ticket still wins over its description's route
        d3 = "**URL/Page:** https://demo.medtricslab.com/gradebook/grouped/\n"
        self.assertEqual(SEL.place_issue(issue("M1-10", description=d3), ROUTES)[1], "known")

    def test_group_route_clicks_get_own_entry(self):
        def pl(tid, clicks):
            d = ("**URL/Page:** https://demo.medtricslab.com/reports/x/\n"
                 + ("Opens a Session > hits View Button\n" if clicks else ""))
            i = issue(tid, description=d)
            s, how = SEL.place_issue(i, ROUTES)
            return SEL._placement(i, s, how, ROUTES)
        placements = [pl("M1-60", True), pl("M1-61", True), pl("M1-62", False), pl("M1-63", False),
                      SEL._placement(issue("M1-64"), ROUTES["surfaces"][1], "module", ROUTES)]
        capture, _, _ = SEL.group(placements, ROUTES)
        by = {tuple(c["tickets"]): c for c in capture}
        self.assertEqual(set(by), {("M1-60",), ("M1-61",), ("M1-62", "M1-63"), ("M1-64",)})
        c = by[("M1-60",)]
        self.assertEqual(c["clicks"], [{"row": True}, {"text": "View"}])
        self.assertEqual((c["url"], c["route_source"], c["write_gated"]), ("/reports/x/", "url_page", ""))
        self.assertEqual(by[("M1-62", "M1-63")]["clicks"], [])
        self.assertEqual(by[("M1-64",)]["clicks"], [])

    def test_capturability(self):
        p = SEL._placement(issue("M1-20"), ROUTES["surfaces"][1], "known", ROUTES)
        self.assertFalse(p["capturable"])
        self.assertEqual(p["reason"], "no demo data")
        p = SEL._placement(issue("M1-32"), ROUTES["surfaces"][2], "module", ROUTES)
        self.assertEqual(p["reason"], "surface not mapped to a demo URL")

    def test_group_multi_persona_and_cap(self):
        s = ROUTES["surfaces"][0]
        placements = [SEL._placement(issue(f"M1-{n}"), s, "rule", ROUTES) for n in range(100, 111)]
        capture, writers, uncap = SEL.group(placements, ROUTES, max_group=8,
                                            sessions={"admin": "valid"})
        self.assertEqual([(c["persona"], len(c["tickets"])) for c in capture],
                         [("admin", 11), ("faculty", 11)])
        self.assertEqual([len(w["tickets"]) for w in writers], [8, 3])
        self.assertTrue(all(w["persona"] == "admin" for w in writers))
        self.assertEqual(capture[0]["session"], "valid")
        self.assertEqual(uncap, [])

    def test_group_write_gated_is_not_shared_across_tickets(self):
        def pl(tid, gated):
            return {"ticket": tid, "persona": "admin", "surface": "Forms > Assignments",
                    "url": "/forms/x/", "capturable": True,
                    "route": {"start": "/forms/x/", "clicks": [], "write_gated": gated, "source": "url_page"}}
        placements = [pl("M1-70", "confirm dialog after Save"), pl("M1-71", ""), pl("M1-72", ""),
                      pl("M1-73", "confirm dialog after Save")]
        capture, _, _ = SEL.group(placements, ROUTES)
        by = {tuple(c["tickets"]): c["write_gated"] for c in capture if c["persona"] == "admin"}
        self.assertEqual(by, {("M1-70", "M1-73"): "confirm dialog after Save", ("M1-71", "M1-72"): ""})

    def test_classifier_bundle_and_merge(self):
        text = SEL.classifier_bundle([issue("M1-40", description="x" * 900)], ROUTES)
        self.assertIn("admin · Forms > Assignments", text)
        self.assertIn("=== M1-40 ===", text)
        placed = {"placements": [], "unplaced": ["M1-40", "M1-41", "M1-42"]}
        answer = [{"ticket": "M1-40", "persona": "admin", "surface": "Gradebook > Charts", "has_ui": True},
                  {"ticket": "M1-41", "has_ui": False},
                  {"ticket": "M1-42", "persona": "admin", "surface": "Invented"}]
        by_id = {t: issue(t) for t in ("M1-40", "M1-41", "M1-42")}
        out = SEL.merge_classifier(placed, answer, ROUTES, by_id)
        self.assertEqual(out["unplaced"], ["M1-42"])
        self.assertEqual({p["ticket"]: p["capturable"] for p in out["placements"]},
                         {"M1-40": True, "M1-41": False})


class Inputs(unittest.TestCase):
    def test_split_ac_quoted_and_added_strings(self):
        desc = 'Add a "Filter to Selected" toggle.\n\nAcceptance criteria:\n- Shows "N selected"'
        diff = ("+<button>Filter to Selected</button>\n"
                "-<button>Old Label</button>\n"
                "+    label=_('Show archived')\n"
                "+ path('forms/picker/', view)\n")
        i = issue("M1-1", description=desc,
                  diffs=[("src/templates/x.html", diff), ("src/apps/forms/urls.py", "+ path('forms/picker/', v)")])
        rec = I.build_input(i)
        self.assertIn("Filter to Selected", rec["description"])
        self.assertNotIn("Acceptance", rec["description"])
        self.assertIn("N selected", rec["acceptance_criteria"])
        self.assertIn("Filter to Selected", rec["ui_strings"])
        self.assertIn("Show archived", rec["ui_strings"])
        self.assertNotIn("Old Label", rec["ui_strings"])
        self.assertEqual(rec["routes"], ["forms/picker/"])
        self.assertEqual(rec["labels"][0], "Filter to Selected")

    def test_labels_user_visible_only(self):
        desc = ('Add "Custom Question" and "DemographicOptions" with "ANSWER_TEXT_TYPES".\n'
                'Email Subject: "Your case log was reviewed"\n'
                'Body: "Hi {name}, Please review"\n')
        diff = ("+<button>Add Question</button>\n"
                "+    label=_('Question text')\n"
                "+    foo = _('some_key')\n")
        i = issue("M1-1", description=desc, diffs=[("src/templates/q.html", diff),
                                                   ("src/apps/q/forms.py",
                                                    "+    q = forms.Field(label='Python Label')\n"
                                                    "+    t = _('Gettext Label')\n")])
        labels = I.build_input(i)["labels"]
        self.assertEqual(labels, ["Custom Question", "Add Question", "Question text", "Gettext Label"])

    def test_route_uses_untrimmed_description(self):
        import kb_route as R
        desc = "word " * 3000 + "\n**URL/Page:** https://demo.medtricslab.com/reports/\n"
        i = issue("M1-1", description=desc)
        self.assertLess(len(I.build_input(i, cap=1000)["description"]), 700)
        self.assertEqual(R.extract(i)["start"], "/reports/")

    def test_cap_and_summary_reuse(self):
        rec = I.build_input(issue("M1-1", description="word " * 5000), cap=1000)
        self.assertLess(len(rec["description"]), 700)
        rec = I.build_input(issue("M1-1", description="word " * 5000),
                            {"CLIENT_TITLE": "T", "EXEC_SUMMARY": "S"}, cap=1000)
        self.assertEqual(rec["client_summary"], "S")
        self.assertLess(len(rec["description"]), 400)

    def test_tests_and_migrations_ignored(self):
        i = issue("M1-1", diffs=[("src/apps/x/tests/test_a.py", "+label='Nope'"),
                                 ("src/apps/x/migrations/0001.py", "+verbose_name='Nope'")])
        self.assertEqual(I.ui_strings(i), [])


class Capture(unittest.TestCase):
    def test_login_detection(self):
        self.assertTrue(C.is_login_page("https://d/accounts/login/?next=/", False, "/accounts/login/"))
        self.assertTrue(C.is_login_page("https://d/x", True, "/accounts/login/"))
        self.assertFalse(C.is_login_page("https://d/administration/", False, "/accounts/login/"))

    def test_box_scale(self):
        self.assertEqual(C.to_image_box({"x": 10, "y": 20.4, "width": 100, "height": 30}, 2),
                         (20, 41, 200, 60))

    def test_resume_plan(self):
        plan = {"run_id": "r", "capture_plan": [{"persona": "admin", "surface": "A"},
                                                {"persona": "admin", "surface": "B"}]}
        self.assertEqual([c["surface"] for c in C.resume_plan(plan, {("admin", "A")})["capture_plan"]],
                         ["B"])

    def test_evidence(self):
        ev = C.shot_evidence("Save", "Save", (1, 2, 3, 4), None, False, "CLEAR", None)
        self.assertEqual((ev["anchor"], ev["label_score"], ev["tenant"]), ("dom_box", 1.0, "CLEAR"))
        self.assertIsNone(C.shot_evidence("Save", "Save draft", (1, 2, 3, 4), None, False, "CLEAR", None))
        self.assertIsNone(C.shot_evidence("Save", "Save", (1, 2, 3, 4), None, True, "CLEAR", None))
        ev = C.shot_evidence("Save", "Save", (1, 2, 3, 4), {"verdict": "FUZZY", "score": 0.9},
                             False, "CLEAR", "REVIEW")
        self.assertEqual((ev["anchor"], ev["tenant"]), ("label_box", "REVIEW"))
        self.assertIsNone(C.shot_evidence("Save", "Save", (1, 2, 3, 4),
                                          {"verdict": "MISS", "score": 0.4}, False, "CLEAR", None))


class Login(unittest.TestCase):
    def test_login_file_and_creds(self):
        with tempfile.TemporaryDirectory() as d:
            f = Path(d, ".env-kb-login")
            self.assertEqual(C.read_login_file(f), {})
            f.write_text("# capture accounts\nKB_ADMIN_EMAIL=a@x.com\nexport KB_ADMIN_PASSWORD='p w'\n"
                         "KB_EMAIL=s@x.com\nKB_PASSWORD=shared\n")
            kv = C.read_login_file(f)
            self.assertEqual(C.creds_for("admin", kv), ("a@x.com", "p w"))
            self.assertEqual(C.creds_for("faculty", kv), ("s@x.com", "shared"))
            self.assertEqual(C.creds_for("admin", {"KB_ADMIN_EMAIL": "a"}), (None, None))

    def test_routes_point_at_capture_site(self):
        import json as _j
        r = _j.loads((Path(C.__file__).parent / "kb_routes.json").read_text())
        self.assertEqual(r["demo"]["base_url"], "https://medtrics.medtricslab.com")
        self.assertEqual(r["demo"]["login_path"], "/users/login/")


class Screen(unittest.TestCase):
    def test_allowlist(self):
        raw = "Welcome to Medtrics Medical Center"
        self.assertEqual(Q.screen_text(raw)["verdict"], "REVIEW")
        self.assertEqual(Q.screen_text(raw, allow_strings=["medtrics medical center"])["verdict"], "CLEAR")

    def test_block_and_course_codes(self):
        self.assertEqual(Q.screen_text("x@kfshrc.edu.sa", "medtrics.sg")["verdict"], "BLOCK")
        r = Q.screen_text("Course MCAP5010 starts")
        self.assertEqual(r["verdict"], "REVIEW")
        self.assertEqual(r["findings"][0]["kind"], "course_codes")

    def test_common_words_are_not_tenant_slugs(self):
        r = Q.screen_text("Schedule · Scheduled · Named · Latest · Informed", "medtricslab.com")
        self.assertEqual(r["verdict"], "CLEAR", r["findings"])
        self.assertEqual(Q.screen_text("lmucvmtest portal")["verdict"], "BLOCK")

    def test_worst(self):
        self.assertEqual(Q.worst("CLEAR", "REVIEW"), "REVIEW")
        self.assertEqual(Q.worst("BLOCK", "CLEAR"), "BLOCK")


@unittest.skipUnless(has_module("PIL"), "Pillow not installed")
class Annotate(unittest.TestCase):
    def test_ring_badge_watermark(self):
        import qa_annotate as A
        from PIL import Image
        im = Image.new("RGB", (1440, 900), "white")
        out = A.annotate(im, A.anchor_from_box((700, 400, 120, 30)), number=3)
        ring_x = 700 + 60 - (120 // 2 + 28)
        self.assertEqual(out.getpixel((ring_x + 1, 415)), A.ACCENT)
        self.assertNotEqual(A.ACCENT, A.RED)
        wm = A.watermark(out.copy())
        self.assertEqual(wm.getpixel((2, 2)), A.RED)
        self.assertEqual(out.size, (1440, 900))   # no caption bar on the KB path

    def test_sidebar_snap(self):
        import qa_annotate as A
        from PIL import Image
        im = Image.new("RGB", (1440, 900), "white")
        out = A.annotate(im, A.anchor_from_box((100, 400, 120, 30)), number=1)
        # inside the badge disc, right of the glyph
        self.assertEqual(out.getpixel((int(0.045 * 1440) + 12, 415)), A.ACCENT)

    def test_ring_stays_inside_frame_at_the_edge(self):
        import qa_annotate as A
        from PIL import Image
        im = Image.new("RGB", (1440, 900), "white")
        out = A.annotate(im, A.anchor_from_box((1330, 200, 110, 34)), number=1)
        self.assertEqual(out.getpixel((1440 - 1 - 5, 217)), A.ACCENT)

    def test_font_is_scalable(self):
        import qa_annotate as A
        from PIL import ImageFont
        self.assertIsInstance(A._font(20), ImageFont.FreeTypeFont)


def prose(ticket="M1-1", shot="s1"):
    g = golden()
    return {"ticket": ticket, "title": g["title"], "steps_heading": "Steps",
            "overview": g["overview"],
            "steps": [{"text": "Click **Filter to Selected**.", "shot": shot, "caption": "The toggle"},
                      {"text": "Click **Save**.", "shot": "missing"}],
            "considerations": g["considerations"], "conclusion": g["conclusion"], "notes": []}


class Spec(unittest.TestCase):
    META = {"ticket": "M1-1", "audience": "Program Administrator", "portal": "Administration portal"}

    def shot(self, frame="f.png", **kw):
        s = {"shot_id": "s1", "ticket": "M1-1", "verified": True, "frame": frame,
             "anchor": "dom_box", "label": "Filter to Selected", "label_score": 1.0,
             "dom_box": [10, 10, 50, 20], "tenant": "CLEAR", "surface": "Forms"}
        s.update(kw)
        return s

    def test_label_as_written(self):
        self.assertEqual(SP.label_as_written("Badge reads **N selected**.", "2 selected"), "N selected")
        self.assertIsNone(SP.label_as_written("Click **Save**.", "Save"))
        self.assertIsNone(SP.label_as_written("Click **Other**.", "Save"))

    def test_unverified_shot_becomes_missing(self):
        spec = SP.build_spec(prose(shot="s9"), [self.shot()], self.META, "no demo data", ".",
                             annotate=False)
        self.assertEqual(spec["steps"][0]["missing_shot"], "screenshot s9 was not verified")
        self.assertEqual(spec["steps"][1]["missing_shot"], "no demo data")

    def test_split_failures(self):
        fails = [V.Failure("G2", "title", "x"), V.Failure("G3", "step 1", "label 'x' is not printed in bold"),
                 V.Failure("G3", "step 2", "label_score missing"), V.Failure("G5", "step 3", "image not found")]
        evidence, prose_f = SP.split_failures(fails)
        self.assertEqual([f.where for f in evidence], ["step 2", "step 3"])
        self.assertEqual([f.gate for f in prose_f], ["G2", "G3"])

    @unittest.skipUnless(has_module("PIL"), "Pillow not installed")
    def test_assemble_annotates_and_passes(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as d:
            frame = Path(d) / "f.png"
            Image.new("RGB", (2880, 1800), "white").save(frame)
            spec, ok, prose_fails = SP.assemble_one(prose(), [self.shot(str(frame))], self.META,
                                                    "no demo data", d)
            self.assertTrue(ok, [str(f) for f in prose_fails])
            self.assertTrue((Path(d) / "assets/M1-1/01.jpg").exists())
            self.assertEqual(spec["steps"][0]["evidence"]["anchor"], "dom_box")
            with Image.open(Path(d) / "assets/M1-1/01.jpg") as im:
                self.assertEqual(im.width, 1200)

    @unittest.skipUnless(has_module("PIL"), "Pillow not installed")
    def test_evidence_failure_demoted_not_retried(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as d:
            frame = Path(d) / "f.png"
            Image.new("RGB", (1440, 900), "white").save(frame)
            spec, ok, prose_fails = SP.assemble_one(prose(), [self.shot(str(frame), anchor="none")],
                                                    self.META, None, d)
            self.assertTrue(ok)
            self.assertEqual(prose_fails, [])
            self.assertIn("missing_shot", spec["steps"][0])

    def test_writer_reply_parsing(self):
        reply = ("noise\n---KB_SPECS_START---\n" + json.dumps({"ticket": "M1-1"}) + "\n" +
                 json.dumps({"ticket": "M1-2"}) + "\n---KB_SPECS_END---\n")
        self.assertEqual([o["ticket"] for o in SP.parse_writer_reply(reply)], ["M1-1", "M1-2"])

    def test_bundle_lists_only_verified_shots(self):
        g = {"group_id": "admin-01", "persona": "admin", "portal": "P", "audience": "A",
             "tickets": ["M1-1", "M1-2"]}
        text = SP.writer_bundle(g, {"M1-1": {"title": "T"}},
                                {"M1-1": [self.shot(), self.shot(shot_id="s2", verified=False)]},
                                {"M1-1": [{"id": "CX-1", "title": "How to X"}]},
                                {"M1-2": "no demo data"})
        self.assertIn("s1 · label", text)
        self.assertNotIn("s2 ·", text)
        self.assertIn("Prior articles (list only, do not merge): CX-1 How to X", text)
        self.assertIn("none (no demo data)", text)

    def test_cmd_bundle_translates_reason_codes_for_the_writer(self):
        import types
        with tempfile.TemporaryDirectory() as d:
            d = Path(d)
            groups = {"writer_groups": [{"group_id": "admin-01", "persona": "admin", "portal": "P",
                                         "audience": "A", "tickets": ["M1-1", "M1-2", "M1-3"]}],
                      "uncapturable": [{"ticket": "M1-1", "reason": "no_screen"},
                                       {"ticket": "M1-2", "reason": "no demo data"}]}
            (d / "g.json").write_text(json.dumps(groups))
            (d / "s.json").write_text(json.dumps({"missing": {"M1-3": {"reason": "needs_write"}}}))
            (d / "i.json").write_text("{}")
            (d / "p.json").write_text("{}")
            a = types.SimpleNamespace(groups=str(d / "g.json"), inputs=str(d / "i.json"),
                                      priors=str(d / "p.json"), shots=str(d / "s.json"),
                                      out_dir=str(d / "out"))
            self.assertEqual(SP.cmd_bundle(a), 0)
            text = (d / "out" / "writer_admin-01.md").read_text()
            self.assertNotIn("no_screen", text)
            self.assertNotIn("needs_write", text)
            self.assertIn("this change has no screen", text)
            self.assertIn("none (no demo data)", text)
            self.assertIn("after you save", text)


@unittest.skipUnless(has_module("docx"), "python-docx not installed")
class Render(unittest.TestCase):
    def test_draft_has_no_provenance_by_default(self):
        import build_kb_doc as B
        from docx import Document
        with tempfile.TemporaryDirectory() as d:
            write_assets(d)
            Path(d, "assets/01.jpg").unlink()
            Path(d, "assets/02.jpg").unlink()
            from PIL import Image
            for n in ("01", "02"):
                Image.new("RGB", (100, 60), "white").save(Path(d, f"assets/{n}.jpg"))
            spec = golden()
            spec["steps"][2]["missing_shot"] = "no demo data"
            B.render(spec, Path(d, "a.docx"), "draft", root=d)
            text = [p.text for p in Document(Path(d, "a.docx")).paragraphs]
            self.assertNotIn("Provenance — strip before publishing", text)
            self.assertIn("[Screenshot missing — no demo data]", text)
            B.render(spec, Path(d, "b.docx"), "draft", root=d, provenance=True)
            text = [p.text for p in Document(Path(d, "b.docx")).paragraphs]
            self.assertIn("Provenance — strip before publishing", text)
            B.render_provenance_md(spec, Path(d, "p.md"))
            self.assertIn("| Step | Source |", Path(d, "p.md").read_text())


@unittest.skipUnless(has_module("pypdf"), "pypdf not installed")
class Pdf(unittest.TestCase):
    def test_verify_reports_missing_text(self):
        import kb_pdf as P
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory() as d:
            w = PdfWriter()
            w.add_blank_page(612, 792)
            w.write(Path(d, "x.pdf"))
            errs, _ = P.verify(Path(d, "x.pdf"), golden())
            self.assertTrue(any("missing text" in e for e in errs))
            self.assertTrue(any("images" in e for e in errs))
            self.assertFalse(any("Letter" in e for e in errs))


class TicketTable(unittest.TestCase):
    def test_every_requested_ticket_has_a_row(self):
        import kb_package as PK
        with tempfile.TemporaryDirectory() as d:
            Path(d, "M1-1").mkdir()
            for e in (".docx", ".md"):
                Path(d, "M1-1", f"CX-DRAFT_How to A{e}").write_text("x")
            rows = [{"ticket": "M1-1", "status": "partial", "files": ["CX-DRAFT_How to A.docx"],
                     "reason": ""},
                    {"ticket": "M1-2", "status": "failed", "files": [], "reason": "G2 title"}]
            groups = {"skipped": [{"ticket": "M1-3", "reason": SEL.NO_CODE},
                                  {"ticket": "M1-4", "reason": "ledger: drafted"}],
                      "uncapturable": [{"ticket": "M1-1", "reason": "no demo data"}]}
            t = {r["ticket"]: r for r in PK.ticket_table(["M1-1", "M1-2", "M1-3", "M1-4", "M1-5"],
                                                         rows, groups, d)}
            self.assertEqual(t["M1-1"]["done"], "done")
            self.assertEqual(t["M1-1"]["files"], "M1-1/CX-DRAFT_How to A (.docx, .md)")
            self.assertIn("no demo data", t["M1-1"]["detail"])
            self.assertEqual((t["M1-2"]["done"], t["M1-2"]["detail"]), ("not done", "error: G2 title"))
            self.assertEqual(t["M1-3"]["detail"], SEL.NO_CODE)
            self.assertEqual(t["M1-4"]["done"], "already done")
            self.assertTrue(t["M1-5"]["detail"].startswith("error"))


class Usage(unittest.TestCase):
    def test_apportion(self):
        with tempfile.TemporaryDirectory() as d:
            Path(d, "writer_admin-01.md").write_text(
                "GROUP\n=== TICKET M1-1 ===\n" + "a" * 300 + "\n=== TICKET M1-2 ===\n" + "b" * 100)
            groups = {"writer_groups": [{"group_id": "admin-01", "tickets": ["M1-1", "M1-2"]}]}
            rows = [{"agent": "kb-writer", "scope": "admin-01", "model": "sonnet", "total_tokens": 4000},
                    {"agent": "kb-writer", "scope": "retry:M1-2", "model": "sonnet", "total_tokens": 500},
                    {"agent": "kb-classifier", "scope": "classify", "model": "haiku", "total_tokens": 200}]
            per = U.apportion(rows, groups, d)
            self.assertEqual(sum(per.values()), 4700)
            self.assertGreater(per["M1-1"], per["M1-2"] - 500)
            self.assertIsNone(U.cost(rows))


if __name__ == "__main__":
    unittest.main()
